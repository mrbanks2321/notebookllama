import os
import re
from typing import List, Optional, Tuple
from pathlib import Path
import pypdf
from docx import Document as DocxDocument
import markdown


class LocalParser:
    """Local parser for various document formats."""
    
    def __init__(self):
        """Initialize the local parser."""
        pass
    
    async def aparse(self, file_path: str) -> "ParsedDocument":
        """Async parse a document."""
        return self.parse(file_path)
    
    def parse(self, file_path: str) -> "ParsedDocument":
        """Parse a document."""
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        if file_path.suffix.lower() == ".pdf":
            return self._parse_pdf(file_path)
        elif file_path.suffix.lower() in [".docx", ".doc"]:
            return self._parse_docx(file_path)
        elif file_path.suffix.lower() in [".md", ".markdown"]:
            return self._parse_markdown(file_path)
        elif file_path.suffix.lower() in [".txt"]:
            return self._parse_text(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_path.suffix}")
    
    def _parse_pdf(self, file_path: Path) -> "ParsedDocument":
        """Parse PDF file."""
        text_content = []
        
        with open(file_path, "rb") as file:
            pdf_reader = pypdf.PdfReader(file)
            
            for page_num, page in enumerate(pdf_reader.pages):
                page_text = page.extract_text()
                if page_text.strip():
                    text_content.append(f"Page {page_num + 1}:\n{page_text}")
        
        return ParsedDocument(
            text="\n\n".join(text_content),
            file_path=str(file_path),
            file_type="pdf"
        )
    
    def _parse_docx(self, file_path: Path) -> "ParsedDocument":
        """Parse DOCX file."""
        doc = DocxDocument(file_path)
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text)
        
        return ParsedDocument(
            text="\n\n".join(text_content),
            file_path=str(file_path),
            file_type="docx"
        )
    
    def _parse_markdown(self, file_path: Path) -> "ParsedDocument":
        """Parse Markdown file."""
        with open(file_path, "r", encoding="utf-8") as file:
            content = file.read()
        
        return ParsedDocument(
            text=content,
            file_path=str(file_path),
            file_type="markdown"
        )
    
    def _parse_text(self, file_path: Path) -> "ParsedDocument":
        """Parse text file."""
        with open(file_path, "r", encoding="utf-8") as file:
            content = file.read()
        
        return ParsedDocument(
            text=content,
            file_path=str(file_path),
            file_type="text"
        )
    
    async def asave_all_images(self, output_dir: str) -> List[str]:
        """Async save all images from the document."""
        # For now, return empty list since we're not extracting images
        # This can be extended later if needed
        return []
    
    async def aget_markdown_documents(self) -> List["MarkdownDocument"]:
        """Async get markdown documents."""
        # For now, return empty list
        # This can be extended later if needed
        return []


class ParsedDocument:
    """Represents a parsed document."""
    
    def __init__(self, text: str, file_path: str, file_type: str):
        """Initialize parsed document."""
        self.text = text
        self.file_path = file_path
        self.file_type = file_type
    
    def get_text(self) -> str:
        """Get the text content."""
        return self.text


class MarkdownDocument:
    """Represents a markdown document."""
    
    def __init__(self, text: str):
        """Initialize markdown document."""
        self.text = text 